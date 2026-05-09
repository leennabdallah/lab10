#!/usr/bin/env python3
"""Build Lab010_Submission.docx with ordered screenshots."""

from pathlib import Path
import shutil
import subprocess
import sys

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    print("Install python-docx: pip install python-docx", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]

STUDENT = "Youssef Al Hajj Youness"

SHOTS = [
    ("submission/screenshots/01_actions_dev_run.png", "GitHub Actions: successful dev workflow (lint/test → build → deploy-dev green)."),
    ("submission/screenshots/02_actions_prod_run_waiting.png", "GitHub Actions: prod deployment waiting on environment approval."),
    ("submission/screenshots/03_prod_approval_dialog.png", "GitHub: Review deployments dialog for prod."),
    ("submission/screenshots/04_kubectl_pods_dev.png", "Terminal: kubectl get pods -n dev -o wide (recent READY/AGE visible)."),
    ("submission/screenshots/05_kubectl_pods_prod.png", "Terminal: kubectl get pods -n prod -o wide (recent READY/AGE visible)."),
    ("submission/screenshots/06_browser_dev.png", "Browser (dev via port-forward): app shows student name + DEV prominently."),
    ("submission/screenshots/07_browser_prod.png", "Browser (prod via port-forward): app showing production traffic."),
    ("submission/screenshots/08_rollout_history.png", "Terminal: kubectl rollout history deployment/backend -n prod showing ≥ 2 revisions."),
    ("submission/screenshots/09_trivy_output.png", "GitHub Actions build logs: Trivy scan summary table."),
    ("submission/screenshots/10_branch_protection.png", "GitHub branch protection rule for main (PR + checks)."),
    ("submission/screenshots/11_terraform_pr_comment.png", "Pull request containing Terraform Plan comment."),
    ("submission/screenshots/12_ci_yml.png", "Screenshot of ci.yml workflow file contents (readable)."),
]


def github_repo_fallback() -> str:
    repo = shutil.which("git")
    if not repo:
        return "https://github.com/<OWNER>/<REPO>"
    try:
        out = subprocess.run(
            ["git", "-C", str(ROOT), "remote", "get-url", "origin"],
            check=True,
            capture_output=True,
            text=True,
        )
        url = out.stdout.strip()
        return url or "https://github.com/<OWNER>/<REPO>"
    except subprocess.CalledProcessError:
        return "https://github.com/<OWNER>/<REPO>"


def main() -> None:
    import os

    repo_url = os.environ.get("LAB010_GITHUB_REPO", github_repo_fallback())

    doc = Document()
    doc.add_heading("Lab 010 CI/CD with GitHub Actions Submission", level=0)
    doc.add_paragraph(f"Student: {STUDENT}")
    doc.add_paragraph(f"Repository: {repo_url}")
    doc.add_paragraph("AWS region: us-east-1 (lab default)")
    doc.add_paragraph("EKS cluster: cicd-lab")
    doc.add_paragraph("Note: Secrets and IAM keys intentionally omitted.")

    doc.add_page_break()

    for rel, caption in SHOTS:
        path = ROOT / rel
        doc.add_heading(Path(rel).name, level=2)
        if path.is_file():
            try:
                doc.add_picture(str(path), width=Inches(6))
            except OSError:
                doc.add_paragraph(f"MISSING SCREENSHOT (read error): {path.name} — capture this before submission.")
        else:
            para = doc.add_paragraph()
            run = para.add_run(f"MISSING SCREENSHOT: {path.name} — capture this before submission.")
            run.bold = True
        doc.add_paragraph(caption)
        doc.add_paragraph()

    out_docx = ROOT / "submission" / "Lab010_Submission.docx"
    doc.save(out_docx)
    print(f"Wrote {out_docx}")

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdf_path = ROOT / "submission" / "Lab010_Submission.pdf"
    if soffice:
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pdf_path.parent), str(out_docx)],
            check=False,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        converted = pdf_path.parent / Path(out_docx.name).with_suffix(".pdf").name
        print(f"Searched LibreOffice conversion output at {converted}")
    else:
        print("LibreOffice not found — convert submission/Lab010_Submission.docx to PDF manually (Word/File → Save as PDF).")


if __name__ == "__main__":
    main()
